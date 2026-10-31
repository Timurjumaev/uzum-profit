"""Excel eksport (B4.10, 03 §8 "eksport endpoint'lari alohida"): jadval qatorlari → formatlangan ``.xlsx``.

Ustun turlari: ``text``, ``int``, ``money`` (so'm, ming ajratgich), ``ratio`` (kasr → foiz formati; qiymat kasr
qoladi), ``date``, ``datetime`` (Toshkent vaqti), ``bool`` (Ha/Yo'q). Pul va nisbat ``Decimal`` dan ``float`` ga
faqat Excel katagiga yozishda aylanadi (hisob-kitob serverda ``Decimal`` bilan bo'lgan); qiymati yo'q katak — bo'sh.
Birinchi qatorlar: jadval nomi va izoh (davr, filtr, yaratilgan vaqt), keyin qotirilgan sarlavha, avtofiltr; ``total``
belgili ustunlarga oxirida ``SUM`` formulasi qo'yiladi.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from upcore.metrics import TASHKENT

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMATS = {
    "int": "#,##0",
    "money": "#,##0",
    "ratio": "0.0%",
    "date": "DD.MM.YYYY",
    "datetime": "DD.MM.YYYY HH:MM",
}
NUMERIC = {"int", "money", "ratio"}
MAX_WIDTH = 60
SHEET_TITLE_MAX = 31  # Excel cheklovi
HEADER_FILL = PatternFill("solid", fgColor="E7ECEA")
FORMULA_PREFIX = ("=", "+", "-", "@")


@dataclass(frozen=True, slots=True)
class Col:
    header: str
    get: Callable[[Any], Any]
    kind: str = "text"
    total: bool = False


def safe_text(value: str) -> str:
    """Excel formula in'ektsiyasi: ``=``, ``+``, ``-``, ``@`` bilan boshlanadigan matn (Uzum nomi, izoh) katakda
    formula bo'lib ishlamasligi uchun apostrof bilan boshlanadi."""
    return "'" + value if value.startswith(FORMULA_PREFIX) else value


def cell_value(value: Any, kind: str) -> Any:
    if value is None:
        return None
    if kind == "bool":
        return "Ha" if value else "Yoʼq"
    if kind in NUMERIC:
        return int(value) if kind == "int" else float(value)
    if kind == "datetime" and isinstance(value, datetime):
        return value.astimezone(TASHKENT).replace(tzinfo=None)
    if kind == "date" and isinstance(value, date):
        return value
    if isinstance(value, (list, tuple, set, frozenset)):
        value = ", ".join(str(v) for v in value)
    if isinstance(value, Decimal):
        value = format(value, "f")
    return safe_text(str(value))


def build(
    title: str,
    columns: Sequence[Col],
    rows: Iterable[Any],
    note: Sequence[str] = (),
    generated: datetime | None = None,
) -> bytes:
    """Bitta varaqli ish daftari baytlari."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title[:SHEET_TITLE_MAX]
    stamp = (generated or datetime.now(TASHKENT)).astimezone(TASHKENT).strftime("%d.%m.%Y %H:%M")
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([" · ".join([*note, f"Yaratildi: {stamp}"])])
    ws["A2"].font = Font(color="666666")
    ws.append([])
    head = 4
    ws.append([c.header for c in columns])
    for n, c in enumerate(columns, 1):
        cell = ws.cell(head, n)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="right" if c.kind in NUMERIC else "left", wrap_text=True
        )
    widths = [len(c.header) for c in columns]
    count = 0
    for row in rows:
        count += 1
        values = [cell_value(c.get(row), c.kind) for c in columns]
        ws.append(values)
        for n, (c, v) in enumerate(zip(columns, values, strict=True)):
            fmt = FORMATS.get(c.kind)
            if fmt:
                ws.cell(head + count, n + 1).number_format = fmt
            if v is not None:
                widths[n] = max(widths[n], len(v) if isinstance(v, str) else 12)
    last = head + count
    if count and any(c.total for c in columns):
        ws.cell(last + 1, 1, "Jami").font = Font(bold=True)
        for n, c in enumerate(columns, 1):
            if c.total:
                letter = get_column_letter(n)
                cell = ws.cell(last + 1, n, f"=SUM({letter}{head + 1}:{letter}{last})")
                cell.number_format = FORMATS[c.kind]
                cell.font = Font(bold=True)
    for n, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(n)].width = min(max(w + 2, 10), MAX_WIDTH)
    ws.freeze_panes = f"B{head + 1}"
    if count:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(columns))}{last}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
