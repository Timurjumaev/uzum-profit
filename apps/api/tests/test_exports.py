"""B4.10 Excel eksport: ``upapi.reports.xlsx`` (birlik) va ``/exports/*.xlsx`` (haqiqiy PostgreSQL).

Eksport ro'yxat endpoint'ini chaqiradi, shuning uchun asosiy tekshiruv — fayldagi qatorlar va yig'indilar ro'yxat
API'si bilan bir xil; rol ruxsatlari, summasiz omborchi va boshqa tashkilot ma'lumotining yo'qligi alohida.
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import test_reports
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from test_reports import OCT, World
from upapi.reports import xlsx
from upapi.reports.xlsx import Col

db = test_reports.db
world = test_reports.world
D = Decimal
STAMP = datetime(2026, 10, 31, 12, 0, tzinfo=UTC)
HEAD = 4  # varaqdagi sarlavha qatori


def sheet(data: bytes) -> Worksheet:
    ws = load_workbook(io.BytesIO(data)).active
    assert ws is not None
    return ws


def table(ws: Worksheet) -> list[list[Any]]:
    """Sarlavhadan keyingi qatorlar (jami qatorisiz)."""
    rows = [list(r) for r in ws.iter_rows(min_row=HEAD + 1, values_only=True)]
    return [r for r in rows if r[0] != "Jami"]


# --- Birlik -----------------------------------------------------------------------------------------


def test_build_formats_and_layout() -> None:
    cols = [
        Col("Nomi", lambda r: r[0]),
        Col("Dona", lambda r: r[1], "int", total=True),
        Col("Summa", lambda r: r[2], "money", total=True),
        Col("Marja", lambda r: r[3], "ratio"),
        Col("Sana", lambda r: r[4], "date"),
        Col("Vaqt", lambda r: r[5], "datetime"),
        Col("Bonus", lambda r: r[6], "bool"),
        Col("Teglar", lambda r: r[7]),
    ]
    rows = [
        (
            "A",
            2,
            D("1500000.50"),
            D("0.125"),
            date(2026, 10, 2),
            datetime(2026, 10, 2, 5, 30, tzinfo=UTC),
            True,
            ["x", "y"],
        ),
        ("B", 3, None, None, None, None, False, []),
    ]
    ws = sheet(xlsx.build("Sotuvlar", cols, rows, ["01.10.2026 – 20.10.2026"], STAMP))
    assert ws.title == "Sotuvlar"
    assert ws["A1"].value == "Sotuvlar"
    assert (
        ws["A2"].value == "01.10.2026 – 20.10.2026 · Yaratildi: 31.10.2026 17:00"
    )  # Toshkent (UTC+5)
    assert [c.value for c in ws[HEAD]] == [c.header for c in cols]
    a, b = ws[HEAD + 1], ws[HEAD + 2]
    assert (a[1].value, a[2].value, a[3].value) == (2, 1500000.5, 0.125)
    assert a[2].number_format == "#,##0" and a[3].number_format == "0.0%"
    assert a[4].value == datetime(2026, 10, 2) and a[4].number_format == "DD.MM.YYYY"
    assert a[5].value == datetime(2026, 10, 2, 10, 30)  # Toshkent vaqti
    assert (a[6].value, b[6].value) == ("Ha", "Yoʼq")
    assert a[7].value == "x, y" and b[7].value in (None, "")
    assert [b[i].value for i in (2, 3, 4, 5)] == [
        None,
        None,
        None,
        None,
    ]  # qiymat yo'q — bo'sh katak
    total = ws[HEAD + 3]
    assert total[0].value == "Jami"
    assert total[1].value == f"=SUM(B{HEAD + 1}:B{HEAD + 2})"
    assert total[3].value is None  # ``total`` belgisiz ustun
    assert ws.freeze_panes == f"B{HEAD + 1}"
    assert ws.auto_filter.ref == f"A{HEAD}:H{HEAD + 2}"


def test_build_empty_has_headers_without_total_or_filter() -> None:
    ws = sheet(xlsx.build("Bo'sh", [Col("A", lambda r: r, "int", total=True)], [], generated=STAMP))
    assert [c.value for c in ws[HEAD]] == ["A"]
    assert ws.max_row == HEAD and not ws.auto_filter.ref


def test_text_cannot_become_formula() -> None:
    cols = [Col("Nomi", lambda r: r)]
    evil = ['=HYPERLINK("http://x")', "+1", "-1", "@SUM(1)", "oddiy"]
    ws = sheet(xlsx.build("T", cols, evil, generated=STAMP))
    values = [r[0] for r in table(ws)]
    assert values[:4] == ["'" + e for e in evil[:4]]
    assert values[4] == "oddiy"
    assert all(c.data_type == "s" for row in ws.iter_rows(min_row=HEAD + 1) for c in row)


def test_sheet_title_is_cut_for_excel() -> None:
    ws = sheet(xlsx.build("x" * 80, [Col("A", lambda r: r)], [], generated=STAMP))
    assert len(ws.title) == 31


# --- Endpoint'lar -------------------------------------------------------------------------------------


def export(w: World, name: str, role: str = "owner", **params: Any) -> Any:
    return w.call(f"exports/{name}.xlsx", role, **params)


def test_sales_export_matches_list(world: World) -> None:
    r = export(world, "sales", **OCT)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == xlsx.XLSX_TYPE
    assert (
        r.headers["content-disposition"] == 'attachment; filename="sotuvlar-20261001-20261020.xlsx"'
    )
    assert r.headers["cache-control"] == "no-store"
    listed = world.get("sales", limit=500, **OCT)["items"]
    assert listed  # bo'sh taqqoslash o'tib ketmasin
    rows = table(sheet(r.content))
    assert len(rows) == len(listed)
    assert [row[1] for row in rows] == [x["order_id"] for x in listed]  # tartib ham bir xil
    revenue = sum(D(x["revenue"]) for x in listed)
    assert sum(D(str(row[8])) for row in rows) == revenue


def test_sales_export_follows_filters_and_hides_other_org(world: World) -> None:
    all_rows = table(sheet(export(world, "sales", **OCT).content))
    b_only = table(sheet(export(world, "sales", shop=str(world.b.shop_id), **OCT).content))
    assert 0 < len(b_only) < len(all_rows)
    listed = world.get("sales", limit=500, shop=str(world.b.shop_id), **OCT)["items"]
    assert len(b_only) == len(listed)
    # Boshqa tashkilotning 5 × 1 000 000 so'mlik sotuvi hech qaysi faylda yo'q.
    assert all(row[8] != 5_000_000 for row in all_rows)


def test_skus_export_matches_list(world: World) -> None:
    r = export(world, "skus", **OCT)
    assert r.status_code == 200, r.text
    listed = world.get("skus", limit=500, **OCT)["items"]
    rows = table(sheet(r.content))
    assert [row[0] for row in rows] == [x["title"] for x in listed]
    assert [D(str(row[14])) for row in rows] == [D(x["net_profit"]) for x in listed]
    bad = export(world, "skus", sort="yoq_ustun", **OCT)
    assert bad.status_code == 422


def test_expenses_export_uzum_and_own(world: World) -> None:
    listed = world.get("expenses", limit=500, **OCT)
    uzum = export(world, "expenses", **OCT)
    assert uzum.status_code == 200, uzum.text
    rows = table(sheet(uzum.content))
    assert len(rows) == len(listed["items"])
    assert sum(D(str(r[4])) for r in rows) == sum(D(x["amount"]) for x in listed["items"])
    own = export(world, "expenses", kind="own", **OCT)
    assert own.status_code == 200, own.text
    assert len(table(sheet(own.content))) == len(listed["own_items"])


def test_purchases_and_inventory_export(world: World) -> None:
    purchases = export(world, "purchases")
    assert purchases.status_code == 200, purchases.text
    listed = world.get("purchases")["items"]
    assert len(table(sheet(purchases.content))) == len(listed)
    stock = export(world, "inventory", as_of="2026-10-20")
    assert stock.status_code == 200, stock.text
    inv = world.get("inventory", as_of="2026-10-20", limit=500)["items"]
    rows = table(sheet(stock.content))
    assert len(rows) == len(inv) > 0
    assert [r[0] for r in rows] == [x["title"] for x in inv]
    assert "ombor-20261020.xlsx" in stock.headers["content-disposition"]


def test_roles(world: World) -> None:
    # Moliya jadvallari: buxgalter va menejer — mumkin; omborchi — yo'q (``export`` ruxsati yo'q).
    for role in ("manager", "accountant"):
        assert export(world, "sales", role, **OCT).status_code == 200
        assert export(world, "skus", role, **OCT).status_code == 200
    for name in ("sales", "skus", "expenses"):
        r = export(world, name, "warehouse", **OCT)
        assert r.status_code == 403, name
        assert r.json()["code"] == "forbidden"
    # Ombor va xaridlar: omborchi ham mumkin (``export_inventory``), lekin summasiz (03 §10).
    owner = sheet(export(world, "inventory", "owner", as_of="2026-10-20").content)
    keeper = sheet(export(world, "inventory", "warehouse", as_of="2026-10-20").content)
    assert "Lot narxi" in [c.value for c in owner[HEAD]]
    assert "Lot narxi" not in [c.value for c in keeper[HEAD]]
    assert "Qoldiq qiymati" not in [c.value for c in keeper[HEAD]]
    own_p = sheet(export(world, "purchases", "owner").content)
    keeper_p = sheet(export(world, "purchases", "warehouse").content)
    assert "Jami (F20)" in [c.value for c in own_p[HEAD]]
    assert "Jami (F20)" not in [c.value for c in keeper_p[HEAD]]


def test_requires_login(world: World) -> None:
    world.api.client.cookies.clear()
    r = world.api.client.get(f"/api/v1/orgs/{world.org_id}/exports/sales.xlsx", params=OCT)
    assert r.status_code in (401, 403)
